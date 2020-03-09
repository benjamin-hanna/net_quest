import os
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime


def edit_esi_form(global_items, filtered_dictionary):

    for dictionary_list in filtered_dictionary['network_info']:

        # global variables
        prospect = global_items[0][1]
        broker = global_items[8][1],
        start_date = global_items[6][1],
        current_med_carrier = global_items[15][1],
        current_pbm = global_items[18][1],
        ees = global_items[13][1],

        # analyses
        is_disruption_analysis = dictionary_list[0]['is_disruption_analysis'],
        is_claims_repricing = dictionary_list[0]['is_claims_repricing'],
        is_questionnaire = dictionary_list[0]['is_questionnaire']


    esi_form = Path(os.environ['HOME'] + '/bin/net_quest/src/assets/intake_forms/esi/CH.ESI_Request Form.docx')
    file_name = "_".join(['CH.ESI_Request Form' , prospect , '.docx' ])
    output_directory = Path(os.environ['HOME'] + '/Desktop/net_quest_output/intake_forms/' + file_name)


    document = Document(esi_form)

    now = datetime.now()
    year = now.strftime("%Y")
    month = now.strftime("%m")
    day = now.strftime("%d")

    array = []
    array.append(day)
    array.append(month)
    array.append(year)

    date = ("/".join(array))

    esi_table = document.tables[0]

    esi_table.cell(0, 1).text = prospect
    esi_table.cell(1, 1).text = broker
    esi_table.cell(2, 1).text = start_date
    # esi_table.cell(3, 1).text = " ".join(current_pbm, current_med_carrier)
    esi_table.cell(4, 1).text = ees
    esi_table.cell(5, 1).text = "See attached plan designs"

    if is_claims_repricing:
        esi_table.cell(9, 1).text = 'Yes'
    else:
        esi_table.cell(9, 1).text = 'No'

    if is_disruption_analysis:
        esi_table.cell(13, 1).text = 'Yes'
    else:
        esi_table.cell(13, 1).text = 'No'

    if is_questionnaire:
        esi_table.cell(15, 1).text = 'Yes'
    else:
        esi_table.cell(15, 1).text = 'No'


    document.save(output_directory)

